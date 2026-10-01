#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""独立验证 Kotlin 侧生成的 OOXML 样本。

═══════════════════════════════════════════════════════════════
 为什么要有这个脚本（而不是在 Kotlin 测试里多写几条断言）
═══════════════════════════════════════════════════════════════

Kotlin 侧的 `OoxmlWriterTest` 全是**字符串包含**断言 —— 它们能证明
"我写了这段文字"，但**证明不了"产出的 XML 是良构的"**。
而良构恰恰是 Excel / Word 能打开文件的前提：一个未闭合的标签会让
文件报"已损坏"，而字符串断言**全绿**。

⇒ 所以这里刻意**换一套实现**去解：

    Kotlin 侧：ZipInputStream + 字符串匹配
    本脚本：  zipfile + ElementTree + **openpyxl（第三方库）**

两套完全独立的实现都认可，才叫"结构真的对了"。
这正是本项目在文件沙箱那边总结过的：
**桩测试验证的是「我们的逻辑自洽」，不是「外面真的能读懂」**。

═══════════════════════════════════════════════════════════════
 用法
═══════════════════════════════════════════════════════════════

    1. 先跑 Kotlin 侧导出样本：
       python tools/verify/run_logic_tests.py
       （`OoxmlWriterTest.导出样本供外部工具验证` 会写样本）
    2. 再跑本脚本：
       python tools/verify/check_ooxml_samples.py

退出码 0 = 全部通过；1 = 有失败项（逐条列出）。
"""

from __future__ import annotations

import os
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

SAMPLE_DIR = os.path.join(tempfile.gettempdir(), "pocketagent-ooxml-samples")

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

_ok: list[str] = []
_fail: list[str] = []


def check(cond: bool, msg: str) -> bool:
    """记录一条结论并**返回它**。

    ⚠️ 必须返回 bool —— 调用点写的是 `if not check(...): return`。
    第一版忘了返回值，于是 `not None` 恒为 True，**每个验证函数都在
    第一行就提前返回**，而脚本照样打印"独立验证通过"、退出码 0。
    表现是"通过 2 项"（只剩两条 `os.path.exists`），
    而 40 多项真实检查**一项都没跑**。

    ⇒ 判据不是"有没有 FAIL"，是"**通过多少项**"。
      项数明显少于预期，就是这种失败。
    """
    (_ok if cond else _fail).append(msg)
    return cond


def check_xml_parts(z: zipfile.ZipFile, label: str) -> None:
    """每个 XML / rels 部件都必须能被标准库解析器读成良构文档。"""
    for name in z.namelist():
        if not (name.endswith(".xml") or name.endswith(".rels")):
            continue
        try:
            ET.fromstring(z.read(name))
            check(True, f"{label} 部件良构：{name}")
        except Exception as exc:  # noqa: BLE001
            check(False, f"{label} 部件**不良构**：{name} -> {exc!r}")


def verify_xlsx() -> None:
    path = os.path.join(SAMPLE_DIR, "sample.xlsx")
    if not check(os.path.exists(path), "xlsx 样本存在"):
        return

    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for need in (
            "[Content_Types].xml",
            "_rels/.rels",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "xl/worksheets/sheet1.xml",
        ):
            check(need in names, f"xlsx 含必需部件 {need}")
        check_xml_parts(z, "xlsx")

    # ★★ 第三方实现（openpyxl）能不能真的读出来
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        check(False, "openpyxl 不可用 —— 缺少最强的那一层验证")
        return

    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        check(False, f"openpyxl 打开失败（Excel 多半也会失败）：{exc!r}")
        return

    check(wb.sheetnames == ["销售"], f"openpyxl 读到工作表名 {wb.sheetnames}")
    ws = wb["销售"]
    check(ws["A1"].value == "商品", f"A1 = {ws['A1'].value!r}")
    check(ws["B1"].value == "数量", f"B1 = {ws['B1'].value!r}")
    check(ws["A2"].value == "苹果", f"A2 = {ws['A2'].value!r}")
    check(ws["B2"].value == 3, f"B2 = {ws['B2'].value!r}（应为数字 3）")
    check(ws["C2"].value == 5.5, f"C2 = {ws['C2'].value!r}")

    # ★ 公式注入免疫：文本 "=SUM(A1)" 必须**原样**是文本，且不是公式类型
    check(ws["A3"].value == "=SUM(A1)", f"A3 原样保留 = {ws['A3'].value!r}")
    check(
        ws["A3"].data_type != "f",
        f"A3 的数据类型不是公式（实际 {ws['A3'].data_type!r}）",
    )


def verify_docx() -> None:
    path = os.path.join(SAMPLE_DIR, "sample.docx")
    if not check(os.path.exists(path), "docx 样本存在"):
        return

    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for need in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
            check(need in names, f"docx 含必需部件 {need}")
        check_xml_parts(z, "docx")
        raw = z.read("word/document.xml")

    try:
        root = ET.fromstring(raw)
    except Exception as exc:  # noqa: BLE001
        check(False, f"docx 正文不良构：{exc!r}")
        return

    texts = [(t.text or "") for t in root.iter(W_NS + "t")]
    check("季度报告" in texts, f"docx 含一级标题（实得文本 {texts}）")
    check("第一段正文。" in texts, "docx 含正文段落")
    check("明细" in texts, "docx 含二级标题")
    check("第二行" in texts, "docx 含换行后的第二行")

    brs = list(root.iter(W_NS + "br"))
    check(len(brs) >= 1, f"docx 有换行标记 w:br（实得 {len(brs)} 个）")

    body = root.find(W_NS + "body")
    check(body is not None, "docx 有 w:body")
    if body is not None and len(body):
        check(
            body[-1].tag == W_NS + "sectPr",
            f"body 的最后一个子元素是 sectPr（实际 {body[-1].tag}）",
        )

    # 加粗与字号：标题必须比正文大
    sizes = []
    for rpr in root.iter(W_NS + "rPr"):
        sz = rpr.find(W_NS + "sz")
        if sz is not None:
            sizes.append(int(sz.get(W_NS + "val")))
    check(len(sizes) >= 2, f"docx 至少有标题与正文两种字号（实得 {sizes}）")
    if len(sizes) >= 2:
        check(max(sizes) > min(sizes), f"标题字号应大于正文（实得 {sizes}）")


def main() -> int:
    if not os.path.isdir(SAMPLE_DIR):
        print(f"找不到样本目录：{SAMPLE_DIR}")
        print("请先跑：python tools/verify/run_logic_tests.py")
        return 1

    verify_xlsx()
    verify_docx()

    print(f"通过 {len(_ok)} 项")
    for msg in _ok:
        print(f"  OK   {msg}")

    if _fail:
        print(f"\n失败 {len(_fail)} 项")
        for msg in _fail:
            print(f"  FAIL {msg}")
        return 1

    print("\n独立验证通过 —— 第三方实现（openpyxl）与标准库解析器都认可了产物结构")
    return 0


if __name__ == "__main__":
    sys.exit(main())
